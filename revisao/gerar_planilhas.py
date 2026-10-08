"""Gera as planilhas cegas da revisão humana dos rótulos, uma por anotador.

Uso (com o ambiente do repositório veritai e o openpyxl instalado nele):
    python revisao/gerar_planilhas.py [--anotadores M I] [--base caminho.sqlite3]

Cada anotador recebe revisao/planilhas/revisao_<iniciais>.xlsx com as abas:
- Afirmações: conjunto-base-v0 (65) + conjunto-validacao-v0 (43), sem o rótulo nem a justificativa da IA;
- Checagens: as 45 de mapeamento/revisao-checagens.csv, sem a classe proposta;
- Vizinhos: amostra de 20 do conjunto-vizinhos-v0 (5 por tipo), sem o rótulo de construção;
- Extração e Faltantes: a amostra de rotulos/extracao-v0, sem pontuação, sinais e alertas.
A ordem das linhas é sorteada por anotador (semente fixa), para que os pares de contraste não
fiquem lado a lado nem na mesma ordem para os dois. O guia está em revisao/README.md, e a
consolidação, em revisao/consolidar.py.
"""

import argparse
import csv
import hashlib
import json
import random
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from veritai.config import caminho_base

RAIZ = Path(__file__).resolve().parents[1]
CLASSES = ["apoiada", "refutada", "insuficiente", "conflitante"]
CLASSES_CHECAGEM = [*CLASSES, "fora_de_escopo"]
CONFIANCA = ["alta", "média", "baixa"]
SIM_NAO = ["s", "n"]
RESPOSTA = PatternFill("solid", fgColor="FFF2CC")  # colunas a preencher


def codigo(identificador: str) -> str:
    """Código opaco do item: o id original ("deflacao-2", "viz-antes_do_release-…") denunciaria o rótulo."""
    return hashlib.sha1(identificador.encode("utf-8")).hexdigest()[:8]


def ler_jsonl(caminho: Path) -> list[dict]:
    return [json.loads(l) for l in caminho.read_text(encoding="utf-8").splitlines() if l.strip()]


def ler_csv(caminho: Path) -> list[dict]:
    with caminho.open(encoding="utf-8-sig") as arquivo:
        return list(csv.DictReader(arquivo))


def evidencias_em_texto(evidencias: list[dict]) -> str:
    partes = []
    for i, e in enumerate(evidencias, start=1):
        cabecalho = f"[{i}] {e.get('fonte', '')} — {e.get('data') or 'sem data'}" + (f" — {e['url']}" if e.get("url") else "")
        partes.append(f"{cabecalho}\n{e['texto']}")
    return "\n\n".join(partes) or "(sem evidência)"


def aba(livro: Workbook, titulo: str, colunas: list[tuple[str, int]], linhas: list[list], respostas: dict[str, list[str] | None]) -> None:
    """Cria uma aba com cabeçalho fixo, larguras, quebra de linha e listas suspensas nas colunas de resposta."""
    folha = livro.create_sheet(titulo)
    folha.append([nome for nome, _ in colunas])
    for linha in linhas:
        folha.append(linha)
    for indice, (nome, largura) in enumerate(colunas, start=1):
        letra = folha.cell(row=1, column=indice).column_letter
        folha.column_dimensions[letra].width = largura
        folha.cell(row=1, column=indice).font = Font(bold=True)
        if nome in respostas:
            for linha in range(1, len(linhas) + 2):
                folha.cell(row=linha, column=indice).fill = RESPOSTA
            if opcoes := respostas[nome]:
                validacao = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
                validacao.error, validacao.errorTitle = f"Use um destes valores: {', '.join(opcoes)}", "Valor inválido"
                folha.add_data_validation(validacao)
                validacao.add(f"{letra}2:{letra}{len(linhas) + 1}")
    for linha in folha.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(wrap_text=True, vertical="top")
    folha.freeze_panes = "A2"


def instrucoes(livro: Workbook, anotador: str) -> None:
    folha = livro.active
    folha.title = "Instruções"
    texto = [
        f"Revisão humana dos rótulos da VeritAI — anotador {anotador}",
        "",
        "Leia o guia completo em veritai-dados/revisao/README.md antes de começar.",
        "Preencha só as colunas amarelas. Não consulte a planilha do outro anotador nem os arquivos de rótulos do repositório.",
        "",
        "Afirmações: julgue pelo que as evidências mostradas dizem, na data da notícia.",
        "  apoiada = a evidência confirma; refutada = a evidência desmente; insuficiente = a evidência não decide;",
        "  conflitante = há evidência a favor e evidência contra.",
        "Checagens: abra o artigo da agência e classifique a alegação checada (fora_de_escopo se não for fato verificável).",
        "Vizinhos: abra o release indicado e julgue considerando só o que estava publicado até a data da notícia.",
        "Extração e Faltantes: siga veritai-dados/rotulos/extracao-v0/README.md.",
        "",
        "Confiança: alta, média ou baixa. Observação: use para dúvidas, problemas na evidência ou no texto.",
    ]
    for linha in texto:
        folha.append([linha])
    folha.column_dimensions["A"].width = 130
    folha["A1"].font = Font(bold=True, size=13)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--anotadores", nargs="+", default=["M", "I"])
    parser.add_argument("--base", type=Path, default=None)
    args = parser.parse_args()

    afirmacoes = [(nome, item) for nome in ("conjunto-base-v0", "conjunto-validacao-v0") for item in ler_jsonl(RAIZ / "rotulos" / f"{nome}.jsonl")]
    with sqlite3.connect(f"file:{args.base or caminho_base()}?mode=ro", uri=True) as conexao:
        por_url = {url: (alegacao, agencia, data) for url, alegacao, agencia, data in conexao.execute("SELECT url, alegacao, agencia, data FROM checagens")}
        url_do_documento = dict(conexao.execute("SELECT id, url FROM documentos").fetchall())
        data_do_documento = dict(conexao.execute("SELECT id, data_publicacao FROM documentos").fetchall())
    checagens = ler_csv(RAIZ / "mapeamento" / "revisao-checagens.csv")
    vizinhos = ler_jsonl(RAIZ / "rotulos" / "conjunto-vizinhos-v0.jsonl")
    sorteio = random.Random(0)
    amostra_vizinhos = [i for tipo in sorted({v["tipo"] for v in vizinhos}) for i in sorteio.sample([v for v in vizinhos if v["tipo"] == tipo], 5)]
    extracao = ler_csv(RAIZ / "rotulos" / "extracao-v0" / "anotacao.csv")
    faltantes = ler_csv(RAIZ / "rotulos" / "extracao-v0" / "faltantes.csv")

    pasta = RAIZ / "revisao" / "planilhas"
    pasta.mkdir(parents=True, exist_ok=True)
    for numero, anotador in enumerate(args.anotadores):
        ordem = random.Random(1000 + numero)  # ordem própria de cada anotador, reproduzível
        livro = Workbook()
        instrucoes(livro, anotador)

        linhas = [[codigo(item["id"]), conjunto, item["data"], item["afirmacao"], evidencias_em_texto(item["evidencias"]), "", "", ""] for conjunto, item in afirmacoes]
        ordem.shuffle(linhas)
        aba(livro, "Afirmações",
            [("codigo", 10), ("conjunto", 20), ("data_noticia", 12), ("afirmacao", 45), ("evidencias", 90), ("classe", 14), ("confianca", 11), ("observacao", 35)],
            linhas, {"classe": CLASSES, "confianca": CONFIANCA, "observacao": None})

        linhas = []
        for c in checagens:
            alegacao, agencia, data = por_url.get(c["url"], ("(alegação não encontrada na base)", "", ""))
            linhas.append([c["url"], agencia, data or "", alegacao, c["veredito_original"], "", "", ""])
        ordem.shuffle(linhas)
        aba(livro, "Checagens",
            [("url", 45), ("agencia", 18), ("data", 12), ("alegacao", 60), ("veredito_original", 18), ("classe", 16), ("confianca", 11), ("observacao", 35)],
            linhas, {"classe": CLASSES_CHECAGEM, "confianca": CONFIANCA, "observacao": None})

        linhas = [[codigo(v["id"]), v["data"], v["afirmacao"], v["titulo_release"], url_do_documento.get(v["documento_base"], ""), data_do_documento.get(v["documento_base"], ""), "", ""] for v in amostra_vizinhos]
        ordem.shuffle(linhas)
        aba(livro, "Vizinhos",
            [("codigo", 10), ("data_noticia", 12), ("afirmacao", 55), ("release_titulo", 45), ("release_url", 45), ("release_publicado_em", 14), ("classe", 14), ("observacao", 35)],
            linhas, {"classe": CLASSES, "observacao": None})

        aba(livro, "Extração",
            [("doc_id", 8), ("titulo", 40), ("url", 40), ("data", 12), ("posicao", 8), ("afirmacao_proposta", 60), ("trecho_origem", 50),
             ("verificavel", 11), ("atomica", 9), ("compreensivel_sozinha", 12), ("afirmacao_corrigida", 50), ("observacao", 30)],
            [[e["doc_id"], e["titulo"], e["url"], e["data"], e["posicao"], e["afirmacao_proposta"], e["trecho_origem"], "", "", "", "", ""] for e in extracao],
            {"verificavel": SIM_NAO, "atomica": SIM_NAO, "compreensivel_sozinha": SIM_NAO, "afirmacao_corrigida": None, "observacao": None})

        aba(livro, "Faltantes",
            [("doc_id", 8), ("titulo", 50), ("url", 50), ("data", 12), ("candidatas_propostas", 12), ("afirmacoes_faltantes", 80)],
            [[f["doc_id"], f["titulo"], f["url"], f["data"], f["candidatas_propostas"], ""] for f in faltantes],
            {"afirmacoes_faltantes": None})

        saida = pasta / f"revisao_{anotador}.xlsx"
        livro.save(saida)
        print(f"{saida.relative_to(RAIZ)}: {len(afirmacoes)} afirmações, {len(checagens)} checagens, {len(amostra_vizinhos)} vizinhos, "
              f"{len(extracao)} candidatas da extração, {len(faltantes)} notícias em Faltantes")


if __name__ == "__main__":
    main()
