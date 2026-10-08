"""Consolida a revisão humana: concordância entre os anotadores, divergências e aplicação nos conjuntos.

Uso (com o ambiente do repositório veritai e o openpyxl instalado nele):
    python revisao/consolidar.py [--anotadores M I]             # passo 1: resumo e divergencias.xlsx
    python revisao/consolidar.py [--anotadores M I] --aplicar   # passo 2: grava as classes nos conjuntos

Passo 1 lê revisao/planilhas/revisao_<iniciais>.xlsx e escreve em revisao/resultado/:
- resumo.md: preenchimento, concordância, kappa de Cohen e matriz entre os anotadores por aba, e
  concordância de cada um com o rótulo proposto (IA ou construção, mostrado só agora);
- divergencias.xlsx: os itens em que os anotadores discordaram, com as observações e o rótulo
  proposto, e a coluna `classe_final` para preencher depois da discussão.

Passo 2 grava a classe acordada (ou a `classe_final` da discussão) em:
- rotulos/conjunto-base-v0.jsonl e conjunto-validacao-v0.jsonl: `rotulo` passa a ser o humano, o
  anterior fica em `rotulo_pre`, `revisao` vira "humana" e `revisores` registra as duas respostas;
- mapeamento/revisao-checagens.csv: `classe_revisada`, `revisor` e `status`; a classe da IA fica em `classe_ia`.
Itens ainda sem classe final continuam "pendente_humano". Vizinhos e Extração só entram no resumo.
"""

import argparse
import csv
import hashlib
import io
import json
import os
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[1]
PARA_CODIGO = {
    "apoiada": "SUPPORTED", "refutada": "REFUTED", "insuficiente": "NOT_ENOUGH_EVIDENCE",
    "conflitante": "CONFLICTING_EVIDENCE", "fora_de_escopo": "FORA_DE_ESCOPO",
}
DE_CODIGO = {v: k for k, v in PARA_CODIGO.items()}
CONJUNTOS = ("conjunto-base-v0", "conjunto-validacao-v0")


def gravar(caminho: Path, conteudo: str) -> None:
    """Grava num temporário e troca no fim: uma falha nunca deixa o arquivo pela metade."""
    temporario = caminho.with_suffix(caminho.suffix + ".tmp")
    temporario.write_text(conteudo, encoding="utf-8")
    os.replace(temporario, caminho)


def codigo(identificador: str) -> str:
    return hashlib.sha1(identificador.encode("utf-8")).hexdigest()[:8]


def ler_aba(caminho: Path, aba: str) -> list[dict]:
    folha = load_workbook(caminho, read_only=True)[aba]
    linhas = folha.iter_rows(values_only=True)
    cabecalho = next(linhas)
    return [{c: ("" if v is None else str(v).strip()) for c, v in zip(cabecalho, linha)} for linha in linhas]


def kappa(pares: list[tuple[str, str]]) -> float | None:
    """Kappa de Cohen entre dois anotadores (classes nominais)."""
    if not pares:
        return None
    n = len(pares)
    observada = sum(a == b for a, b in pares) / n
    ca, cb = Counter(a for a, _ in pares), Counter(b for _, b in pares)
    esperada = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / n**2
    return 1.0 if esperada == 1 else (observada - esperada) / (1 - esperada)


def comparar(nome: str, respostas: dict[str, dict[str, str]], propostos: dict[str, str], a: str, b: str, linhas: list[str]) -> list[str]:
    """Escreve no resumo a comparação de uma aba; devolve as chaves com divergência."""
    chaves = sorted(propostos)
    preenchidas = {x: [k for k in chaves if respostas[x].get(k)] for x in (a, b)}
    pares = [(respostas[a][k], respostas[b][k]) for k in chaves if respostas[a].get(k) and respostas[b].get(k)]
    k = kappa(pares)
    linhas.append(f"\n## {nome}\n")
    linhas.append(f"- Itens: {len(chaves)}; preenchidos por {a}: {len(preenchidas[a])}; por {b}: {len(preenchidas[b])}; pelos dois: {len(pares)}.")
    if pares:
        linhas.append(f"- Concordância entre {a} e {b}: {sum(x == y for x, y in pares) / len(pares):.1%}; kappa de Cohen: {k:.2f}.")
        for x in (a, b):
            com_proposto = [respostas[x][c] == propostos[c] for c in preenchidas[x]]
            linhas.append(f"- {x} concorda com o rótulo proposto em {sum(com_proposto) / len(com_proposto):.1%} dos itens que preencheu.")
        classes = sorted({c for par in pares for c in par})
        linhas.append(f"\n| {a} \\\\ {b} | " + " | ".join(classes) + " |")
        linhas.append("|---" * (len(classes) + 1) + "|")
        for ca in classes:
            linhas.append(f"| {ca} | " + " | ".join(str(sum(1 for x, y in pares if x == ca and y == cb)) for cb in classes) + " |")
    return [c for c in chaves if respostas[a].get(c) and respostas[b].get(c) and respostas[a][c] != respostas[b][c]]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--anotadores", nargs=2, default=["M", "I"])
    parser.add_argument("--aplicar", action="store_true")
    args = parser.parse_args()
    a, b = args.anotadores
    planilhas = {x: RAIZ / "revisao" / "planilhas" / f"revisao_{x}.xlsx" for x in (a, b)}
    resultado = RAIZ / "revisao" / "resultado"
    resultado.mkdir(parents=True, exist_ok=True)

    itens = {nome: [json.loads(l) for l in (RAIZ / "rotulos" / f"{nome}.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()] for nome in CONJUNTOS}
    por_codigo = {codigo(i["id"]): (nome, i) for nome, lista in itens.items() for i in lista}
    with (RAIZ / "mapeamento" / "revisao-checagens.csv").open(encoding="utf-8") as arquivo:
        checagens = list(csv.DictReader(arquivo))
    if malformadas := [c["url"] for c in checagens if None in c]:
        raise SystemExit(f"revisao-checagens.csv tem linhas com colunas a mais (vírgula sem aspas?): {malformadas}. Corrija antes de consolidar.")
    vizinhos = {codigo(v["id"]): v for v in (json.loads(l) for l in (RAIZ / "rotulos" / "conjunto-vizinhos-v0.jsonl").read_text(encoding="utf-8").splitlines() if l)}

    respostas = {aba: {x: {} for x in (a, b)} for aba in ("Afirmações", "Checagens", "Vizinhos")}
    detalhes = {x: {} for x in (a, b)}
    for x, caminho in planilhas.items():
        for linha in ler_aba(caminho, "Afirmações"):
            respostas["Afirmações"][x][linha["codigo"]] = PARA_CODIGO.get(linha["classe"], "")
            detalhes[x][linha["codigo"]] = (linha["confianca"], linha["observacao"])
        for linha in ler_aba(caminho, "Checagens"):
            respostas["Checagens"][x][linha["url"]] = PARA_CODIGO.get(linha["classe"], "")
            detalhes[x][linha["url"]] = (linha["confianca"], linha["observacao"])
        for linha in ler_aba(caminho, "Vizinhos"):
            respostas["Vizinhos"][x][linha["codigo"]] = PARA_CODIGO.get(linha["classe"], "")
            detalhes[x][linha["codigo"]] = ("", linha["observacao"])

    linhas = [f"# Revisão humana: resumo ({a} × {b})", "", "Gerado por `revisao/consolidar.py`. O rótulo proposto é o da IA (afirmações, checagens) ou o da construção (vizinhos)."]
    propostos_afirmacoes = {c: i["rotulo"] for c, (_, i) in por_codigo.items()}
    propostos_checagens = {c["url"]: c["classe_revisada"] for c in checagens}
    amostra_vizinhos = {c: vizinhos[c]["rotulo"] for c in respostas["Vizinhos"][a] if c in vizinhos}
    divergencias = {
        "Afirmações": comparar("Afirmações (desenvolvimento + validação)", respostas["Afirmações"], propostos_afirmacoes, a, b, linhas),
        "Checagens": comparar("Checagens", respostas["Checagens"], propostos_checagens, a, b, linhas),
        "Vizinhos": comparar("Vizinhos (amostra; rótulo de construção)", respostas["Vizinhos"], amostra_vizinhos, a, b, linhas),
    }

    # Extração: precisão por anotador e concordância por pergunta; cobertura com as faltantes.
    linhas.append("\n## Extração (amostra de 30 notícias)\n")
    extracao = {x: ler_aba(planilhas[x], "Extração") for x in (a, b)}
    faltas = {x: ler_aba(planilhas[x], "Faltantes") for x in (a, b)}
    for pergunta in ("verificavel", "atomica", "compreensivel_sozinha"):
        pares = [(p[pergunta], q[pergunta]) for p, q in zip(extracao[a], extracao[b]) if p[pergunta] and q[pergunta]]
        k = kappa(pares)
        linhas.append(f"- `{pergunta}`: respondida pelos dois em {len(pares)} de {len(extracao[a])}; concordância {sum(p == q for p, q in pares) / len(pares):.1%}; kappa {k:.2f}." if pares else f"- `{pergunta}`: sem respostas dos dois.")
    for x in (a, b):
        respondidas = [l for l in extracao[x] if l["verificavel"]]
        verificaveis = sum(l["verificavel"] == "s" for l in respondidas)
        faltantes = sum(len([t for t in l["afirmacoes_faltantes"].splitlines() if t.strip() and t.strip().lower() != "nenhuma"]) for l in faltas[x])
        if respondidas:
            boas = [l for l in respondidas if l["verificavel"] == "s"]
            completas = sum(l["atomica"] == "s" and l["compreensivel_sozinha"] == "s" for l in boas)
            cobertura = verificaveis / (verificaveis + faltantes) if verificaveis + faltantes else 0
            linhas.append(
                f"- {x}: precisão (verificáveis) {verificaveis / len(respondidas):.1%}; atômicas e compreensíveis sozinhas {completas / len(boas):.1%} das verificáveis; "
                f"{faltantes} faltante(s); cobertura {cobertura:.1%}." if boas else f"- {x}: nenhuma candidata marcada como verificável."
            )
    (resultado / "resumo.md").write_text("\n".join(linhas) + "\n", encoding="utf-8")
    print("\n".join(linhas))

    caminho_div = resultado / "divergencias.xlsx"
    if not args.aplicar:
        livro = Workbook()
        livro.remove(livro.active)
        textos = {c: i["afirmacao"] for c, (_, i) in por_codigo.items()} | {c["url"]: c["url"] for c in checagens} | {c: v["afirmacao"] for c, v in vizinhos.items()}
        propostos = propostos_afirmacoes | propostos_checagens | amostra_vizinhos
        for aba, chaves in divergencias.items():
            folha = livro.create_sheet(aba)
            folha.append(["chave", "afirmacao_ou_url", a, f"confianca_{a}", f"observacao_{a}", b, f"confianca_{b}", f"observacao_{b}", "proposto", "classe_final"])
            for c in chaves:
                folha.append([c, textos.get(c, ""), DE_CODIGO.get(respostas[aba][a][c], ""), *detalhes[a][c], DE_CODIGO.get(respostas[aba][b][c], ""), *detalhes[b][c],
                              DE_CODIGO.get(propostos.get(c, ""), propostos.get(c, "")), ""])
            validacao = DataValidation(type="list", formula1='"' + ",".join(PARA_CODIGO) + '"', allow_blank=True)
            folha.add_data_validation(validacao)
            validacao.add(f"J2:J{max(len(chaves), 1) + 1}")
            for coluna, largura in zip("ABCDEFGHIJ", (12, 60, 14, 10, 30, 14, 10, 30, 14, 16)):
                folha.column_dimensions[coluna].width = largura
            for celula in folha[1]:
                celula.font = Font(bold=True)
            for linha in folha.iter_rows(min_row=2):
                for celula in linha:
                    celula.alignment = Alignment(wrap_text=True, vertical="top")
                linha[-1].fill = PatternFill("solid", fgColor="FFF2CC")
        livro.save(caminho_div)
        print(f"\nDivergências para discutir em {caminho_div.relative_to(RAIZ)}; depois, rode com --aplicar.")
        return

    finais: dict[str, str] = {}
    if caminho_div.exists():
        for aba in ("Afirmações", "Checagens"):
            for linha in ler_aba(caminho_div, aba):
                if linha["classe_final"] in PARA_CODIGO:
                    finais[linha["chave"]] = PARA_CODIGO[linha["classe_final"]]

    def final(aba: str, chave: str) -> str | None:
        ra, rb = respostas[aba][a].get(chave), respostas[aba][b].get(chave)
        if ra and ra == rb:
            return ra
        return finais.get(chave)

    aplicados, novos = Counter(), {}
    for nome, lista in itens.items():
        for item in lista:
            chave = codigo(item["id"])
            if classe := final("Afirmações", chave):
                item.setdefault("rotulo_pre", item["rotulo"])
                item["rotulo"], item["revisao"] = classe, "humana"
                item["revisores"] = {a: respostas["Afirmações"][a].get(chave, ""), b: respostas["Afirmações"][b].get(chave, "")}
                aplicados[nome] += 1
        novos[RAIZ / "rotulos" / f"{nome}.jsonl"] = "".join(json.dumps(i, ensure_ascii=False) + "\n" for i in lista)
    campos = list(checagens[0]) + (["classe_ia"] if "classe_ia" not in checagens[0] else [])
    for c in checagens:
        if classe := final("Checagens", c["url"]):
            c.setdefault("classe_ia", c["classe_revisada"])
            if not c["classe_ia"]:
                c["classe_ia"] = c["classe_revisada"]
            c["classe_revisada"], c["revisor"], c["status"] = classe, f"{a}+{b}", "revisado"
            aplicados["checagens"] += 1
        c.setdefault("classe_ia", "")
    saida = io.StringIO()
    escritor = csv.DictWriter(saida, fieldnames=campos, lineterminator="\n")
    escritor.writeheader()
    escritor.writerows(checagens)
    novos[RAIZ / "mapeamento" / "revisao-checagens.csv"] = saida.getvalue()
    # Só grava depois de montar tudo: um erro antes daqui não toca em nenhum arquivo.
    for caminho, conteudo in novos.items():
        gravar(caminho, conteudo)
    print(f"\nAplicado: {dict(aplicados)}. Itens sem classe final continuam pendentes.")


if __name__ == "__main__":
    main()
